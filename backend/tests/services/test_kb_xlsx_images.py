import io
import zipfile
from io import BytesIO
from types import SimpleNamespace
from xml.etree import ElementTree as ET
from zipfile import ZIP_DEFLATED, ZipFile

import pytest
from openpyxl import Workbook
from openpyxl.drawing.image import Image as WorksheetImage
from openpyxl.drawing.spreadsheet_drawing import (
    AbsoluteAnchor,
    AnchorMarker,
    OneCellAnchor,
    TwoCellAnchor,
)
from openpyxl.drawing.xdr import XDRPoint2D, XDRPositiveSize2D
from PIL import Image as PILImage

from app.services import kb_xlsx_images as xlsx_module
from app.services.kb_xlsx_images import (
    XlsxImageExtractionError,
    extract_xlsx_images,
)


def _png_bytes(color: str = "red") -> bytes:
    output = BytesIO()
    PILImage.new("RGB", (2, 3), color=color).save(output, format="PNG")
    return output.getvalue()


def _workbook_with_images(*, image_count: int = 1) -> bytes:
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "Quarterly Results"
    for index in range(image_count):
        worksheet.add_image(
            WorksheetImage(BytesIO(_png_bytes())), f"{chr(66 + index)}{index + 3}"
        )
    output = BytesIO()
    workbook.save(output)
    return output.getvalue()


def _rewrite_package(content: bytes, transform) -> bytes:
    source = ZipFile(BytesIO(content))
    output = BytesIO()
    with source, ZipFile(output, "w", ZIP_DEFLATED) as target:
        for info in source.infolist():
            payload = source.read(info.filename)
            payload = transform(info.filename, payload)
            target.writestr(info.filename, payload)
        for name, payload in getattr(transform, "extra_entries", ()):
            target.writestr(name, payload)
    return output.getvalue()


def test_extract_xlsx_images_returns_payload_and_cell_position():
    extracted = extract_xlsx_images(_workbook_with_images())

    assert len(extracted) == 1
    image = extracted[0]
    assert image.sheet_name == "Quarterly Results"
    assert image.anchor == "B3"
    assert (image.row, image.column) == (2, 1)
    assert image.content_type == "image/png"
    assert image.data == _png_bytes()
    assert image.end_anchor is None


def test_extract_xlsx_images_preserves_two_cell_anchor_and_offsets():
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "Anchored"
    image = WorksheetImage(BytesIO(_png_bytes()))
    image.anchor = TwoCellAnchor(
        _from=AnchorMarker(col=4, row=6, colOff=360, rowOff=720),
        to=AnchorMarker(col=8, row=10, colOff=180, rowOff=240),
    )
    worksheet.add_image(image)
    output = BytesIO()
    workbook.save(output)

    record = extract_xlsx_images(output.getvalue())[0]

    assert record.anchor == "E7"
    assert record.end_anchor == "I11"
    assert (record.row_offset_emu, record.column_offset_emu) == (720, 360)
    assert (record.end_row_offset_emu, record.end_column_offset_emu) == (240, 180)


def test_extract_xlsx_images_rejects_image_count_over_limit():
    with pytest.raises(XlsxImageExtractionError, match="more images"):
        extract_xlsx_images(_workbook_with_images(image_count=2), max_count=1)


def test_extract_xlsx_images_rejects_image_payload_over_limit():
    with pytest.raises(XlsxImageExtractionError, match="byte limit"):
        extract_xlsx_images(_workbook_with_images(), max_image_bytes=1)


def test_extract_xlsx_images_rejects_external_relationships():
    def add_external_relationship(name: str, payload: bytes) -> bytes:
        if name == "_rels/.rels":
            payload = payload.replace(
                b"</Relationships>",
                b'<Relationship Id="rIdExternal" '
                b'Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink" '
                b'Target="https://example.invalid/" TargetMode="External"/>'
                b"</Relationships>",
            )
        return payload

    with pytest.raises(XlsxImageExtractionError, match="External"):
        extract_xlsx_images(
            _rewrite_package(_workbook_with_images(), add_external_relationship)
        )


def test_extract_xlsx_images_rejects_macro_parts():
    def add_macro_project(name: str, payload: bytes) -> bytes:
        return payload

    add_macro_project.extra_entries = (("xl/vbaProject.bin", b"macro"),)
    with pytest.raises(XlsxImageExtractionError, match="Macro"):
        extract_xlsx_images(
            _rewrite_package(_workbook_with_images(), add_macro_project)
        )


def test_extract_xlsx_images_rejects_unsupported_embedded_format():
    def replace_image(name: str, payload: bytes) -> bytes:
        return (
            b"<svg xmlns='http://www.w3.org/2000/svg'/>"
            if name == "xl/media/image1.png"
            else payload
        )

    with pytest.raises(XlsxImageExtractionError, match="Invalid or unsafe"):
        extract_xlsx_images(_rewrite_package(_workbook_with_images(), replace_image))


def test_extract_xlsx_images_rejects_absolute_anchors():
    workbook = Workbook()
    worksheet = workbook.active
    image = WorksheetImage(BytesIO(_png_bytes()))
    image.anchor = AbsoluteAnchor(
        pos=XDRPoint2D(x=0, y=0),
        ext=XDRPositiveSize2D(cx=914400, cy=914400),
    )
    worksheet.add_image(image)
    output = BytesIO()
    workbook.save(output)

    with pytest.raises(XlsxImageExtractionError, match="Absolute"):
        extract_xlsx_images(output.getvalue())


def test_extract_xlsx_images_rejects_compressed_input_over_limit():
    with pytest.raises(XlsxImageExtractionError, match="compressed size"):
        extract_xlsx_images(b"0" * (64 * 1024 * 1024 + 1))


def _info(
    name, *, size=1, compress_size=1, flag_bits=0, compress_type=zipfile.ZIP_STORED
):
    return SimpleNamespace(
        filename=name,
        file_size=size,
        compress_size=compress_size,
        flag_bits=flag_bits,
        compress_type=compress_type,
    )


class _Archive:
    def __init__(self, infos, payloads=None, error=None):
        self._infos = infos
        self._payloads = payloads or {}
        self._error = error
        self.read_payload = b""

    def infolist(self):
        return list(self._infos)

    def open(self, info, _mode):
        if self._error:
            raise self._error
        payload = self._payloads.get(info.filename, b"x" * info.file_size)
        return io.BytesIO(payload)

    def read(self, _info):
        return self.read_payload


def test_validate_limits_rejects_bad_types_and_values(monkeypatch):
    for content in ("text", b""):
        with pytest.raises(XlsxImageExtractionError):
            xlsx_module._validate_limits(content, 1, 1)
    monkeypatch.setattr(xlsx_module, "MAX_COMPRESSED_ARCHIVE_BYTES", 1)
    with pytest.raises(XlsxImageExtractionError, match="compressed size"):
        xlsx_module._validate_limits(b"xx", 1, 1)
    monkeypatch.setattr(xlsx_module, "MAX_COMPRESSED_ARCHIVE_BYTES", 64 * 1024 * 1024)
    for count in (True, "1", -1):
        with pytest.raises(XlsxImageExtractionError, match="max_count"):
            xlsx_module._validate_limits(b"x", count, 1)
    for size in (True, "1", -1):
        with pytest.raises(XlsxImageExtractionError, match="max_image_bytes"):
            xlsx_module._validate_limits(b"x", 1, size)


def test_validate_archive_rejects_unsafe_entries_and_limits(monkeypatch):
    required = [
        _info("[Content_Types].xml"),
        _info("_rels/.rels"),
        _info("xl/workbook.xml"),
    ]
    with pytest.raises(XlsxImageExtractionError, match="too many"):
        monkeypatch.setattr(xlsx_module, "MAX_ARCHIVE_ENTRIES", 1)
        xlsx_module._validate_archive(_Archive(required))

    monkeypatch.setattr(xlsx_module, "MAX_ARCHIVE_ENTRIES", 10_000)
    for name in ("x" * 1025, "", "a\\b", "/abs", "../evil"):
        with pytest.raises(XlsxImageExtractionError):
            xlsx_module._validate_archive(_Archive(required + [_info(name)]))
    with pytest.raises(XlsxImageExtractionError):
        xlsx_module._validate_archive(
            _Archive(required + [_info("[Content_Types].xml")])
        )

    with pytest.raises(XlsxImageExtractionError, match="Encrypted"):
        xlsx_module._validate_archive(
            _Archive(required + [_info("xl/a.xml", flag_bits=1)])
        )
    with pytest.raises(XlsxImageExtractionError, match="compression"):
        xlsx_module._validate_archive(
            _Archive(required + [_info("xl/a.xml", compress_type=zipfile.ZIP_BZIP2)])
        )
    with pytest.raises(XlsxImageExtractionError, match="invalid archive size"):
        xlsx_module._validate_archive(_Archive(required + [_info("xl/a.xml", size=-1)]))
    with pytest.raises(XlsxImageExtractionError, match="uncompressed size"):
        monkeypatch.setattr(xlsx_module, "MAX_UNCOMPRESSED_ARCHIVE_BYTES", 1)
        xlsx_module._validate_archive(_Archive(required))


def test_validate_archive_checks_required_parts_and_inflated_sizes(monkeypatch):
    monkeypatch.setattr(xlsx_module, "MAX_UNCOMPRESSED_ARCHIVE_BYTES", 10**9)
    with pytest.raises(XlsxImageExtractionError, match="package structure"):
        xlsx_module._validate_archive(_Archive([_info("[Content_Types].xml")]))

    required = [
        _info("[Content_Types].xml", size=2),
        _info("_rels/.rels"),
        _info("xl/workbook.xml"),
    ]
    with pytest.raises(XlsxImageExtractionError, match="size mismatch"):
        xlsx_module._validate_archive(_Archive(required, {"[Content_Types].xml": b"x"}))
    with pytest.raises(XlsxImageExtractionError, match="Invalid XLSX ZIP"):
        xlsx_module._validate_archive(_Archive(required, error=OSError("boom")))
    monkeypatch.setattr(xlsx_module, "MAX_UNCOMPRESSED_ARCHIVE_BYTES", 3)
    with pytest.raises(XlsxImageExtractionError, match="uncompressed size"):
        xlsx_module._validate_archive(_Archive(required))


def test_safe_xml_part_rejects_missing_oversized_and_invalid_parts(monkeypatch):
    with pytest.raises(XlsxImageExtractionError, match="missing"):
        xlsx_module._safe_xml_part(_Archive([]), {}, "a.xml")

    monkeypatch.setattr(xlsx_module, "MAX_XML_PART_BYTES", 1)
    with pytest.raises(XlsxImageExtractionError, match="exceeds"):
        xlsx_module._safe_xml_part(
            _Archive([]), {"a.xml": _info("a.xml", size=2)}, "a.xml"
        )

    monkeypatch.setattr(xlsx_module, "MAX_XML_PART_BYTES", 1024)
    with pytest.raises(XlsxImageExtractionError, match="unsafe"):
        xlsx_module._safe_xml_part(
            _Archive([], {"a.xml": b"<"}), {"a.xml": _info("a.xml")}, "a.xml"
        )


def test_relationship_and_target_validation(monkeypatch):
    with pytest.raises(XlsxImageExtractionError, match="target"):
        xlsx_module._resolve_target("xl/workbook.xml", "")
    with pytest.raises(XlsxImageExtractionError, match="target"):
        xlsx_module._resolve_target("xl/workbook.xml", "a\\b")
    with pytest.raises(XlsxImageExtractionError, match="escapes"):
        xlsx_module._resolve_target("xl/workbook.xml", "../../evil")

    rel = ET.fromstring(
        '<Relationships><Relationship Id="rId1" Type="http://x/image" '
        'Target="media/a.png"/></Relationships>'
    )
    monkeypatch.setattr(xlsx_module, "_safe_xml_part", lambda *a, **k: rel)
    entries = {"xl/_rels/workbook.xml.rels": _info("x")}
    assert xlsx_module._read_relationships(None, entries, "xl/workbook.xml") == {
        "rId1": ("xl/media/a.png", "http://x/image")
    }


@pytest.mark.parametrize(
    "xml",
    [
        '<Relationships><Relationship Type="t" Target="a"/></Relationships>',
        '<Relationships><Relationship Id="r" Target="a"/></Relationships>',
        '<Relationships><Relationship Id="r" Type="t"/></Relationships>',
        '<Relationships><Relationship Id="r" Type="t" Target="a"/>'
        '<Relationship Id="r" Type="t" Target="a"/></Relationships>',
        '<Relationships><Relationship Id="r" Type="t" Target="a" '
        'TargetMode="External"/></Relationships>',
        '<Relationships><Relationship Id="r" Type="http://x/externalLink" '
        'Target="a"/></Relationships>',
    ],
)
def test_read_relationships_rejects_invalid_entries(monkeypatch, xml):
    monkeypatch.setattr(
        xlsx_module, "_safe_xml_part", lambda *a, **k: ET.fromstring(xml)
    )
    with pytest.raises(XlsxImageExtractionError):
        xlsx_module._read_relationships(
            None, {"xl/_rels/workbook.xml.rels": _info("x")}, "xl/workbook.xml"
        )


def _drawing_xml(*anchors: str) -> ET.Element:
    return ET.fromstring(
        f'<xdr:wsDr xmlns:xdr="{xlsx_module._XDR_NS}" '
        f'xmlns:a="{xlsx_module._DRAWING_NS}" '
        f'xmlns:r="{xlsx_module._OFFICE_REL_NS}">' + "".join(anchors) + "</xdr:wsDr>"
    )


def _anchor_xml(*blips: str) -> str:
    return (
        "<xdr:oneCellAnchor><xdr:from><xdr:col>1</xdr:col>"
        "<xdr:colOff>0</xdr:colOff><xdr:row>2</xdr:row>"
        "<xdr:rowOff>0</xdr:rowOff></xdr:from>"
        "<xdr:pic><xdr:blipFill>" + "".join(blips) + "</xdr:blipFill></xdr:pic>"
        "</xdr:oneCellAnchor>"
    )


def test_embedded_blips_validate_picture_references():
    assert (
        xlsx_module._embedded_blip_ids(
            _drawing_xml(
                "<xdr:oneCellAnchor><xdr:from><xdr:col>0</xdr:col>"
                "<xdr:row>0</xdr:row></xdr:from></xdr:oneCellAnchor>"
            )
        )
        == []
    )

    with pytest.raises(XlsxImageExtractionError, match="Linked"):
        xlsx_module._embedded_blip_ids(
            _drawing_xml(_anchor_xml('<a:blip r:link="rId1"/>'))
        )

    with pytest.raises(XlsxImageExtractionError, match="not embedded"):
        xlsx_module._embedded_blip_ids(_drawing_xml(_anchor_xml("<a:blip/>")))

    blips = xlsx_module._embedded_blip_ids(
        _drawing_xml(_anchor_xml('<a:blip r:embed="rId1"/>'))
    )
    assert len(blips) == 1
    assert blips[0][0].embed == "rId1"
    assert blips[0][1]._from.col == 1
    assert blips[0][1]._from.row == 2


def test_embedded_blips_collect_repeated_grouped_pictures():
    root = _drawing_xml(
        _anchor_xml(
            '<a:blip r:embed="rId1"/>',
            '<xdr:grpSp><xdr:grpSpPr/><a:blip r:embed="rId2"/></xdr:grpSp>',
        )
    )
    blips = xlsx_module._embedded_blip_ids(root)
    assert [blip.embed for blip, _anchor in blips] == ["rId1", "rId2"]
    assert blips[0][1] is blips[1][1]


def test_validate_image_rejects_unsupported_animated_and_corrupt(monkeypatch):
    bmp = io.BytesIO()
    PILImage.new("RGB", (2, 2)).save(bmp, format="BMP")
    with pytest.raises(XlsxImageExtractionError, match="Unsupported"):
        xlsx_module._validate_image(bmp.getvalue())

    with pytest.raises(XlsxImageExtractionError, match="Invalid or unsafe"):
        xlsx_module._validate_image(b"not an image")

    frames = [PILImage.new("RGB", (2, 2), color) for color in ("red", "blue")]
    gif = io.BytesIO()
    frames[0].save(gif, format="GIF", save_all=True, append_images=frames[1:])
    with pytest.raises(XlsxImageExtractionError, match="Animated"):
        xlsx_module._validate_image(gif.getvalue())

    monkeypatch.setattr(xlsx_module, "MAX_IMAGE_DIMENSION", 0)
    with pytest.raises(XlsxImageExtractionError, match="dimensions"):
        xlsx_module._validate_image(_png_bytes())


def test_build_image_record_validates_anchor_types_and_coordinates():
    from openpyxl.drawing.spreadsheet_drawing import AbsoluteAnchor

    with pytest.raises(XlsxImageExtractionError, match="Absolute"):
        xlsx_module._build_image_record("S", AbsoluteAnchor(), "image/png", b"x")
    with pytest.raises(XlsxImageExtractionError, match="Unsupported XLSX image anchor"):
        xlsx_module._build_image_record("S", object(), "image/png", b"x")

    marker = AnchorMarker(col=1, row=2, colOff=3, rowOff=4)
    one = OneCellAnchor(_from=marker)
    record = xlsx_module._build_image_record("S", one, "image/png", b"x")
    assert record.anchor == "B3"
    assert record.end_anchor is None

    two = TwoCellAnchor(_from=marker, to=marker)
    assert (
        xlsx_module._build_image_record("S", two, "image/png", b"x").end_anchor == "B3"
    )

    for marker in (
        SimpleNamespace(row=True, col=0, rowOff=0, colOff=0),
        SimpleNamespace(row=-1, col=0, rowOff=0, colOff=0),
        SimpleNamespace(row=0, col=0, rowOff=-1, colOff=0),
    ):
        with pytest.raises(XlsxImageExtractionError, match="anchor"):
            xlsx_module._marker_values(marker)


def test_extract_rejects_non_zip_and_macro_enabled_content_types():
    with pytest.raises(XlsxImageExtractionError, match="Invalid XLSX ZIP"):
        extract_xlsx_images(b"not a zip")

    def add_macro_enabled(name: str, payload: bytes) -> bytes:
        if name == "[Content_Types].xml":
            payload = payload.replace(b"<Types", b'<Types foo="macroEnabled"', 1)
        return payload

    with pytest.raises(XlsxImageExtractionError, match="Macro-enabled"):
        extract_xlsx_images(
            _rewrite_package(_workbook_with_images(), add_macro_enabled)
        )


def test_relationship_part_size_limit(monkeypatch):
    monkeypatch.setattr(xlsx_module, "MAX_XML_PART_BYTES", 0)
    with pytest.raises(XlsxImageExtractionError, match="relationship part"):
        xlsx_module._reject_external_relationships(
            _Archive([]), {"_rels/.rels": _info("_rels/.rels", size=1)}
        )


def _workbook_xml(sheets: str = "") -> ET.Element:
    return ET.fromstring(
        f'<workbook xmlns="{xlsx_module._WORKBOOK_NS}" '
        f'xmlns:r="{xlsx_module._OFFICE_REL_NS}"><sheets>{sheets}</sheets></workbook>'
    )


def test_worksheet_drawings_rejects_missing_excessive_and_invalid_sheets(
    monkeypatch,
):
    monkeypatch.setattr(
        xlsx_module, "_safe_xml_part", lambda *a, **k: ET.fromstring("<workbook/>")
    )
    with pytest.raises(XlsxImageExtractionError, match="worksheet list"):
        xlsx_module._worksheet_drawings(_Archive([]), {})

    monkeypatch.setattr(
        xlsx_module, "_safe_xml_part", lambda *a, **k: _workbook_xml("<sheet/>")
    )
    monkeypatch.setattr(xlsx_module, "_read_relationships", lambda *a, **k: {})
    with pytest.raises(XlsxImageExtractionError, match="worksheet relationship"):
        xlsx_module._worksheet_drawings(_Archive([]), {})

    monkeypatch.setattr(xlsx_module, "MAX_WORKSHEETS", 0)
    with pytest.raises(XlsxImageExtractionError, match="worksheet list"):
        xlsx_module._worksheet_drawings(_Archive([]), {})


def test_worksheet_drawings_skips_non_worksheets_and_missing_rels(monkeypatch):
    sheet = '<sheet name="S" r:id="rId1"/>'
    monkeypatch.setattr(
        xlsx_module, "_safe_xml_part", lambda *a, **k: _workbook_xml(sheet)
    )
    monkeypatch.setattr(
        xlsx_module,
        "_read_relationships",
        lambda *a, **k: {"rId1": ("xl/chartsheets/sheet1.xml", "http://x/chartsheet")},
    )
    assert xlsx_module._worksheet_drawings(_Archive([]), {}) == []

    monkeypatch.setattr(
        xlsx_module,
        "_read_relationships",
        lambda *a, **k: {"rId1": ("xl/worksheets/sheet1.xml", "http://x/worksheet")},
    )
    assert xlsx_module._worksheet_drawings(_Archive([]), {}) == []


def test_worksheet_drawings_skips_other_rel_types_and_enforces_count(monkeypatch):
    sheet = '<sheet name="S" r:id="rId1"/>'

    def worksheet_xml(*args, **_kwargs):
        if args[2] == "xl/workbook.xml":
            return _workbook_xml(sheet)
        return ET.fromstring(
            f'<worksheet xmlns="{xlsx_module._WORKBOOK_NS}" '
            f'xmlns:r="{xlsx_module._OFFICE_REL_NS}">'
            '<drawing r:id="rId2"/><drawing r:id="rId3"/></worksheet>'
        )

    monkeypatch.setattr(xlsx_module, "_safe_xml_part", worksheet_xml)
    entries = {"xl/worksheets/_rels/sheet1.xml.rels": _info("x")}
    monkeypatch.setattr(
        xlsx_module,
        "_read_relationships",
        lambda *a, **k: (
            {"rId1": ("xl/worksheets/sheet1.xml", "http://x/worksheet")}
            if a[2] == "xl/workbook.xml"
            else {"rId2": ("xl/drawings/drawing1.xml", "http://x/image")}
        ),
    )
    assert xlsx_module._worksheet_drawings(_Archive([]), entries) == []

    monkeypatch.setattr(
        xlsx_module,
        "_read_relationships",
        lambda *a, **k: (
            {"rId1": ("xl/worksheets/sheet1.xml", "http://x/worksheet")}
            if a[2] == "xl/workbook.xml"
            else {
                "rId2": ("xl/drawings/drawing1.xml", "http://x/drawing"),
                "rId3": ("xl/drawings/drawing2.xml", "http://x/drawing"),
            }
        ),
    )
    monkeypatch.setattr(xlsx_module, "MAX_WORKSHEETS", 1)
    with pytest.raises(XlsxImageExtractionError, match="too many worksheet drawings"):
        xlsx_module._worksheet_drawings(_Archive([]), entries)


def test_worksheet_drawings_ignores_stale_drawing_relationship(monkeypatch):
    sheet = '<sheet name="S" r:id="rId1"/>'

    def worksheet_xml(*args, **_kwargs):
        if args[2] == "xl/workbook.xml":
            return _workbook_xml(sheet)
        return ET.fromstring(
            f'<worksheet xmlns="{xlsx_module._WORKBOOK_NS}" '
            f'xmlns:r="{xlsx_module._OFFICE_REL_NS}"/>'
        )

    monkeypatch.setattr(xlsx_module, "_safe_xml_part", worksheet_xml)
    entries = {"xl/worksheets/_rels/sheet1.xml.rels": _info("x")}
    monkeypatch.setattr(
        xlsx_module,
        "_read_relationships",
        lambda *a, **k: (
            {"rId1": ("xl/worksheets/sheet1.xml", "http://x/worksheet")}
            if a[2] == "xl/workbook.xml"
            else {"rId2": ("xl/drawings/drawing1.xml", "http://x/drawing")}
        ),
    )
    assert xlsx_module._worksheet_drawings(_Archive([]), entries) == []


def test_extract_rejects_drawing_attached_to_multiple_sheets(monkeypatch):
    monkeypatch.setattr(
        xlsx_module, "_safe_xml_part", lambda *a, **k: ET.fromstring("<xdr/>")
    )
    monkeypatch.setattr(xlsx_module, "_read_relationships", lambda *a, **k: {})
    monkeypatch.setattr(xlsx_module, "_embedded_blip_ids", lambda _drawing: [])
    entries = {"xl/drawings/drawing1.xml": _info("x")}
    with pytest.raises(XlsxImageExtractionError, match="multiple worksheets"):
        xlsx_module._extract_drawn_images(
            _Archive([]),
            entries,
            [("A", "xl/drawings/drawing1.xml"), ("B", "xl/drawings/drawing1.xml")],
            max_count=1,
            max_image_bytes=1,
        )


def test_extract_drawn_images_rejects_unanchored_missing_and_mismatched_images(
    monkeypatch,
):
    blip = SimpleNamespace(embed="rId1")
    anchor = SimpleNamespace()
    monkeypatch.setattr(
        xlsx_module, "_safe_xml_part", lambda *a, **k: ET.fromstring("<xdr/>")
    )
    monkeypatch.setattr(
        xlsx_module, "_embedded_blip_ids", lambda _drawing: [(blip, anchor)]
    )
    entries = {"xl/drawings/drawing1.xml": _info("x")}
    worksheets = [("S", "xl/drawings/drawing1.xml")]

    monkeypatch.setattr(xlsx_module, "_read_relationships", lambda *a, **k: {})
    with pytest.raises(XlsxImageExtractionError, match="unanchored"):
        xlsx_module._extract_drawn_images(
            _Archive([]), entries, worksheets, max_count=1, max_image_bytes=100
        )

    monkeypatch.setattr(
        xlsx_module,
        "_read_relationships",
        lambda *a, **k: {"rId1": ("xl/media/a.png", xlsx_module._IMAGE_REL_TYPE)},
    )
    with pytest.raises(XlsxImageExtractionError, match="image is missing"):
        xlsx_module._extract_drawn_images(
            _Archive([]), entries, worksheets, max_count=1, max_image_bytes=100
        )

    entries["xl/media/a.png"] = _info("xl/media/a.png", size=2)
    with pytest.raises(XlsxImageExtractionError, match="size mismatch"):
        xlsx_module._extract_drawn_images(
            _Archive([]),
            entries,
            worksheets,
            max_count=1,
            max_image_bytes=100,
        )
