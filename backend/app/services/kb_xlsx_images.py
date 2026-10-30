"""Bounded extraction of worksheet-embedded raster images from XLSX packages."""

from __future__ import annotations

import posixpath
import warnings
import zipfile
from dataclasses import dataclass
from io import BytesIO

from pathlib import PurePosixPath
from defusedxml import ElementTree as SafeET
from openpyxl.drawing.spreadsheet_drawing import (
    AbsoluteAnchor,
    OneCellAnchor,
    SpreadsheetDrawing,
    TwoCellAnchor,
)
from openpyxl.utils import get_column_letter
from PIL import Image, UnidentifiedImageError


DEFAULT_MAX_IMAGES = 8
DEFAULT_MAX_IMAGE_BYTES = 10 * 1024 * 1024
MAX_COMPRESSED_ARCHIVE_BYTES = 64 * 1024 * 1024
MAX_UNCOMPRESSED_ARCHIVE_BYTES = 128 * 1024 * 1024
MAX_ARCHIVE_ENTRIES = 10_000
MAX_WORKSHEETS = 512
MAX_XML_PART_BYTES = 4 * 1024 * 1024
MAX_IMAGE_PIXELS = 40_000_000
MAX_IMAGE_DIMENSION = 32_768
MAX_ANCHOR_OFFSET_EMU = 100_000_000

_WORKBOOK_NS = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"
_OFFICE_REL_NS = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
_IMAGE_REL_TYPE = (
    "http://schemas.openxmlformats.org/officeDocument/2006/relationships/image"
)
_SUPPORTED_IMAGE_MIME_TYPES = {
    "PNG": "image/png",
    "JPEG": "image/jpeg",
    "GIF": "image/gif",
    "WEBP": "image/webp",
}


class XlsxImageExtractionError(ValueError):
    """An invalid, unsupported, or over-budget XLSX image package."""


@dataclass(frozen=True, slots=True)
class XlsxImage:
    """An image and its worksheet anchor, with zero-based row and column."""

    sheet_name: str
    anchor: str
    row: int
    column: int
    row_offset_emu: int
    column_offset_emu: int
    content_type: str
    data: bytes
    end_anchor: str | None = None
    end_row: int | None = None
    end_column: int | None = None
    end_row_offset_emu: int | None = None
    end_column_offset_emu: int | None = None


def extract_xlsx_images(
    content: bytes,
    *,
    max_count: int = DEFAULT_MAX_IMAGES,
    max_image_bytes: int = DEFAULT_MAX_IMAGE_BYTES,
) -> list[XlsxImage]:
    """Extract worksheet raster images without loading worksheet cell data.

    ``max_image_bytes`` bounds the total bytes returned for embedded worksheet
    images. The package is separately capped at 64 MiB compressed, 128 MiB
    uncompressed, and 10,000 ZIP entries before image or drawing parts are read.
    External relationships, macro-enabled packages, unsupported images, and
    unsupported image anchors raise :class:`XlsxImageExtractionError`.
    """
    _validate_limits(content, max_count, max_image_bytes)

    try:
        if not zipfile.is_zipfile(BytesIO(content)):
            raise XlsxImageExtractionError("Invalid XLSX ZIP archive")
        with zipfile.ZipFile(BytesIO(content)) as archive:
            entries = _validate_archive(archive)
            _reject_external_relationships(archive, entries)
            _reject_macros(entries, archive)
            worksheets = _worksheet_drawings(archive, entries)
            return _extract_drawn_images(
                archive,
                entries,
                worksheets,
                max_count=max_count,
                max_image_bytes=max_image_bytes,
            )
    except XlsxImageExtractionError:
        raise
    except (OSError, RuntimeError, ValueError, KeyError, zipfile.BadZipFile) as exc:
        raise XlsxImageExtractionError(
            "Invalid or unsupported XLSX image package"
        ) from exc


def _validate_limits(content: bytes, max_count: int, max_image_bytes: int) -> None:
    if not isinstance(content, bytes):
        raise XlsxImageExtractionError("XLSX content must be bytes")
    if not content:
        raise XlsxImageExtractionError("XLSX content is empty")
    if len(content) > MAX_COMPRESSED_ARCHIVE_BYTES:
        raise XlsxImageExtractionError("XLSX compressed size exceeds the limit")
    if isinstance(max_count, bool) or not isinstance(max_count, int) or max_count < 0:
        raise XlsxImageExtractionError("max_count must be a non-negative integer")
    if (
        isinstance(max_image_bytes, bool)
        or not isinstance(max_image_bytes, int)
        or max_image_bytes < 0
    ):
        raise XlsxImageExtractionError("max_image_bytes must be a non-negative integer")


def _validate_archive(archive: zipfile.ZipFile) -> dict[str, zipfile.ZipInfo]:
    infos = archive.infolist()
    if len(infos) > MAX_ARCHIVE_ENTRIES:
        raise XlsxImageExtractionError("XLSX contains too many archive entries")

    entries: dict[str, zipfile.ZipInfo] = {}
    declared_size = 0
    for info in infos:
        name = info.filename
        if len(name) > 1024:
            raise XlsxImageExtractionError("XLSX archive path exceeds the limit")
        path = PurePosixPath(name)
        if (
            not name
            or "\\" in name
            or path.is_absolute()
            or ".." in path.parts
            or name in entries
        ):
            raise XlsxImageExtractionError(
                "XLSX contains an unsafe or duplicate archive path"
            )
        if info.flag_bits & 0x1:
            raise XlsxImageExtractionError("Encrypted XLSX entries are not supported")
        if info.compress_type not in (zipfile.ZIP_STORED, zipfile.ZIP_DEFLATED):
            raise XlsxImageExtractionError("Unsupported XLSX archive compression")
        if info.file_size < 0 or info.compress_size < 0:
            raise XlsxImageExtractionError("XLSX contains an invalid archive size")
        declared_size += info.file_size
        if declared_size > MAX_UNCOMPRESSED_ARCHIVE_BYTES:
            raise XlsxImageExtractionError("XLSX uncompressed size exceeds the limit")
        entries[name] = info

    required_parts = {"[Content_Types].xml", "_rels/.rels", "xl/workbook.xml"}
    if not required_parts.issubset(entries):
        raise XlsxImageExtractionError("Invalid XLSX package structure")

    # Check actual inflated bytes, CRCs, and declared-size consistency before
    # openpyxl's XML parsers or image decoders see the package.
    actual_size = 0
    try:
        for info in infos:
            with archive.open(info, "r") as stream:
                entry_size = 0
                while block := stream.read(64 * 1024):
                    entry_size += len(block)
                    actual_size += len(block)
                    if (
                        entry_size > info.file_size
                        or actual_size > MAX_UNCOMPRESSED_ARCHIVE_BYTES
                    ):
                        raise XlsxImageExtractionError(
                            "XLSX uncompressed size exceeds the limit"
                        )
                if entry_size != info.file_size:
                    raise XlsxImageExtractionError("XLSX archive entry size mismatch")
    except XlsxImageExtractionError:
        raise
    except (OSError, RuntimeError, zipfile.BadZipFile) as exc:
        raise XlsxImageExtractionError("Invalid XLSX ZIP archive") from exc
    return entries


def _safe_xml_part(
    archive: zipfile.ZipFile, entries: dict[str, zipfile.ZipInfo], path: str
):
    info = entries.get(path)
    if info is None:
        raise XlsxImageExtractionError("XLSX relationship target is missing")
    if info.file_size > MAX_XML_PART_BYTES:
        raise XlsxImageExtractionError("XLSX XML part exceeds the limit")
    try:
        return SafeET.fromstring(archive.read(info))
    except Exception as exc:
        raise XlsxImageExtractionError("Invalid or unsafe XLSX XML part") from exc


def _reject_external_relationships(
    archive: zipfile.ZipFile, entries: dict[str, zipfile.ZipInfo]
) -> None:
    for path, info in entries.items():
        if not path.endswith(".rels"):
            continue
        if info.file_size > MAX_XML_PART_BYTES:
            raise XlsxImageExtractionError("XLSX relationship part exceeds the limit")
        root = _safe_xml_part(archive, entries, path)
        for rel in root:
            if rel.attrib.get("TargetMode", "").lower() == "external":
                raise XlsxImageExtractionError(
                    "External XLSX relationships are not supported"
                )


def _reject_macros(
    entries: dict[str, zipfile.ZipInfo], archive: zipfile.ZipFile
) -> None:
    if any(
        name.lower().startswith(("xl/externallinks/", "xl/activex/"))
        or name.lower().endswith(("vbaproject.bin", "vbadata.xml"))
        for name in entries
    ):
        raise XlsxImageExtractionError(
            "Macro or external-link XLSX packages are not supported"
        )
    content_types = _safe_xml_part(archive, entries, "[Content_Types].xml")
    if any(
        "macroenabled" in value.lower() or "vbaproject" in value.lower()
        for element in content_types.iter()
        for value in element.attrib.values()
    ):
        raise XlsxImageExtractionError("Macro-enabled XLSX packages are not supported")


def _relationship_part(source_part: str) -> str:
    directory, filename = posixpath.split(source_part)
    return posixpath.join(directory, "_rels", f"{filename}.rels")


def _resolve_target(source_part: str, target: str) -> str:
    if not target or "\\" in target:
        raise XlsxImageExtractionError("Invalid XLSX relationship target")
    target_path = posixpath.normpath(
        target.lstrip("/")
        if target.startswith("/")
        else posixpath.join(posixpath.dirname(source_part), target)
    )
    if target_path in ("", ".", "..") or target_path.startswith("../"):
        raise XlsxImageExtractionError("XLSX relationship target escapes the package")
    return target_path


def _read_relationships(
    archive: zipfile.ZipFile,
    entries: dict[str, zipfile.ZipInfo],
    part_path: str,
) -> dict[str, tuple[str, str]]:
    rels_path = _relationship_part(part_path)
    root = _safe_xml_part(archive, entries, rels_path)
    result: dict[str, tuple[str, str]] = {}
    for rel in root:
        rel_id = rel.attrib.get("Id")
        target = rel.attrib.get("Target")
        rel_type = rel.attrib.get("Type")
        if not rel_id or not target or not rel_type or rel_id in result:
            raise XlsxImageExtractionError("Invalid XLSX relationship")
        if rel.attrib.get("TargetMode", "").lower() == "external":
            raise XlsxImageExtractionError(
                "External XLSX relationships are not supported"
            )
        if rel_type.endswith("/externalLink"):
            raise XlsxImageExtractionError("External XLSX links are not supported")
        result[rel_id] = (_resolve_target(part_path, target), rel_type)
    return result


def _worksheet_drawings(
    archive: zipfile.ZipFile, entries: dict[str, zipfile.ZipInfo]
) -> list[tuple[str, str]]:
    workbook = _safe_xml_part(archive, entries, "xl/workbook.xml")
    workbook_rels = _read_relationships(archive, entries, "xl/workbook.xml")
    sheets = workbook.find(f"{{{_WORKBOOK_NS}}}sheets")
    if sheets is None or len(sheets) > MAX_WORKSHEETS:
        raise XlsxImageExtractionError("Invalid or excessive XLSX worksheet list")

    drawings: list[tuple[str, str]] = []
    sheet_names: set[str] = set()
    for sheet in sheets:
        name = sheet.attrib.get("name")
        rel_id = sheet.attrib.get(f"{{{_OFFICE_REL_NS}}}id")
        if not name or name in sheet_names or not rel_id or rel_id not in workbook_rels:
            raise XlsxImageExtractionError("Invalid XLSX worksheet relationship")
        sheet_names.add(name)
        worksheet_path, rel_type = workbook_rels[rel_id]
        if not rel_type.endswith("/worksheet"):
            continue
        sheet_rels_path = _relationship_part(worksheet_path)
        if sheet_rels_path not in entries:
            continue
        sheet_rels = _read_relationships(archive, entries, worksheet_path)
        for drawing_path, drawing_type in sheet_rels.values():
            if drawing_type.endswith("/drawing"):
                drawings.append((name, drawing_path))
                if len(drawings) > MAX_WORKSHEETS:
                    raise XlsxImageExtractionError(
                        "XLSX contains too many worksheet drawings"
                    )
    return drawings


def _extract_drawn_images(
    archive: zipfile.ZipFile,
    entries: dict[str, zipfile.ZipInfo],
    worksheets: list[tuple[str, str]],
    *,
    max_count: int,
    max_image_bytes: int,
) -> list[XlsxImage]:
    extracted: list[XlsxImage] = []
    total_image_bytes = 0
    seen_drawings: set[str] = set()

    for sheet_name, drawing_path in worksheets:
        if drawing_path in seen_drawings:
            raise XlsxImageExtractionError(
                "A drawing is attached to multiple worksheets"
            )
        seen_drawings.add(drawing_path)
        drawing_root = _safe_xml_part(archive, entries, drawing_path)
        try:
            drawing = SpreadsheetDrawing.from_tree(drawing_root)
        except Exception as exc:
            raise XlsxImageExtractionError(
                "Unsupported XLSX worksheet drawing"
            ) from exc
        relationships = _read_relationships(archive, entries, drawing_path)
        embedded_blips = _embedded_blip_ids(drawing)
        embedded_rel_ids = {blip.embed for blip, _anchor in embedded_blips}
        image_rel_ids = {
            rel_id
            for rel_id, (_, rel_type) in relationships.items()
            if rel_type == _IMAGE_REL_TYPE
        }
        if image_rel_ids != embedded_rel_ids:
            raise XlsxImageExtractionError(
                "Unsupported or unanchored XLSX drawing image"
            )

        for blip, anchor in embedded_blips:
            if len(extracted) >= max_count:
                raise XlsxImageExtractionError("XLSX contains more images than allowed")
            rel = relationships.get(blip.embed)
            if rel is None or rel[1] != _IMAGE_REL_TYPE:
                raise XlsxImageExtractionError("Unsupported XLSX image relationship")
            image_path = rel[0]
            image_info = entries.get(image_path)
            if image_info is None:
                raise XlsxImageExtractionError("XLSX embedded image is missing")
            if image_info.file_size > max_image_bytes - total_image_bytes:
                raise XlsxImageExtractionError(
                    "XLSX embedded images exceed the byte limit"
                )
            try:
                data = archive.read(image_info)
            except (OSError, RuntimeError, zipfile.BadZipFile) as exc:
                raise XlsxImageExtractionError(
                    "Unable to read XLSX embedded image"
                ) from exc
            if len(data) != image_info.file_size:
                raise XlsxImageExtractionError("XLSX embedded image size mismatch")
            total_image_bytes += len(data)
            mime_type = _validate_image(data)
            extracted.append(_build_image_record(sheet_name, anchor, mime_type, data))

    return extracted


def _embedded_blip_ids(drawing: SpreadsheetDrawing):
    blips = []
    anchors = (
        list(drawing.absoluteAnchor)
        + list(drawing.oneCellAnchor)
        + list(drawing.twoCellAnchor)
    )
    for anchor in anchors:
        picture = anchor.pic or (anchor.groupShape and anchor.groupShape.pic)
        if picture is None or picture.blipFill is None:
            continue
        blip = picture.blipFill.blip
        if blip is None:
            raise XlsxImageExtractionError(
                "XLSX drawing picture has no image reference"
            )
        if blip.link:
            raise XlsxImageExtractionError("Linked XLSX images are not supported")
        if not blip.embed:
            raise XlsxImageExtractionError("XLSX drawing picture is not embedded")
        blips.append((blip, anchor))
    return blips


def _validate_image(data: bytes) -> str:
    try:
        with warnings.catch_warnings():
            warnings.simplefilter("error", Image.DecompressionBombWarning)
            with Image.open(BytesIO(data)) as image:
                mime_type = _SUPPORTED_IMAGE_MIME_TYPES.get(image.format or "")
                if mime_type is None:
                    raise XlsxImageExtractionError(
                        "Unsupported XLSX embedded image format"
                    )
                width, height = image.size
                if (
                    width <= 0
                    or height <= 0
                    or width > MAX_IMAGE_DIMENSION
                    or height > MAX_IMAGE_DIMENSION
                    or width * height > MAX_IMAGE_PIXELS
                ):
                    raise XlsxImageExtractionError(
                        "XLSX embedded image dimensions exceed the limit"
                    )
                try:
                    image.seek(1)
                except EOFError:
                    pass
                else:
                    raise XlsxImageExtractionError(
                        "Animated XLSX embedded images are not supported"
                    )
                image.verify()
                return mime_type
    except XlsxImageExtractionError:
        raise
    except (
        OSError,
        ValueError,
        Image.DecompressionBombError,
        Image.DecompressionBombWarning,
        UnidentifiedImageError,
    ) as exc:
        raise XlsxImageExtractionError("Invalid or unsafe XLSX embedded image") from exc


def _build_image_record(
    sheet_name: str, anchor: object, content_type: str, data: bytes
) -> XlsxImage:
    if isinstance(anchor, AbsoluteAnchor):
        raise XlsxImageExtractionError("Absolute XLSX image anchors are not supported")
    if isinstance(anchor, OneCellAnchor):
        start = anchor._from
        end = None
    elif isinstance(anchor, TwoCellAnchor):
        start = anchor._from
        end = anchor.to
    else:
        raise XlsxImageExtractionError("Unsupported XLSX image anchor")

    row, column, row_offset, column_offset = _marker_values(start)
    end_values = _marker_values(end) if end is not None else None
    return XlsxImage(
        sheet_name=sheet_name,
        anchor=f"{get_column_letter(column + 1)}{row + 1}",
        row=row,
        column=column,
        row_offset_emu=row_offset,
        column_offset_emu=column_offset,
        content_type=content_type,
        data=data,
        end_anchor=(
            f"{get_column_letter(end_values[1] + 1)}{end_values[0] + 1}"
            if end_values is not None
            else None
        ),
        end_row=end_values[0] if end_values is not None else None,
        end_column=end_values[1] if end_values is not None else None,
        end_row_offset_emu=end_values[2] if end_values is not None else None,
        end_column_offset_emu=end_values[3] if end_values is not None else None,
    )


def _marker_values(marker: object) -> tuple[int, int, int, int]:
    values = (marker.row, marker.col, marker.rowOff, marker.colOff)
    if any(isinstance(value, bool) or not isinstance(value, int) for value in values):
        raise XlsxImageExtractionError("Invalid XLSX image anchor coordinates")
    row, column, row_offset, column_offset = values
    if not (0 <= row < 1_048_576 and 0 <= column < 16_384):
        raise XlsxImageExtractionError("XLSX image anchor is outside worksheet bounds")
    if not (
        0 <= row_offset <= MAX_ANCHOR_OFFSET_EMU
        and 0 <= column_offset <= MAX_ANCHOR_OFFSET_EMU
    ):
        raise XlsxImageExtractionError(
            "XLSX image anchor offset is outside supported bounds"
        )
    return row, column, row_offset, column_offset
